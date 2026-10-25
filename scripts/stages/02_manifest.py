"""建立分析清单：DDR 纳入/排除规则、分割掩膜按分级划分重新分配、DeepDRiD 眼级标签与排除。
输出 results/02_manifest/images.tsv（每行一张分级图像）与 exclusions.tsv、summary.json。"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from scripts.lib.repro import ROOT, load_config, write_provenance  # noqa: E402

ROLE = {"train": "ddr_train", "valid": "ddr_dev", "test": "ddr_test"}
ORDER = {"ddr_train": 0, "ddr_dev": 1, "ddr_test": 2}
NEAR_DUP_BITS = 2


def hamming(a: str, b: str) -> int:
    return bin(int(a, 16) ^ int(b, 16)).count("1")


def near_dup_pairs(df: pd.DataFrame) -> list:
    """dHash 汉明距离 ≤ NEAR_DUP_BITS 的图像对（按 16 位前缀分桶加速，四段任一相同即比较）。"""
    keys, hs = df.image_key.to_numpy(), df.dhash.to_numpy()
    pairs = set()
    for seg in range(4):
        buckets = {}
        for i, h in enumerate(hs):
            buckets.setdefault(h[seg * 16:(seg + 1) * 16], []).append(i)
        for idx in buckets.values():
            for a in range(len(idx)):
                for b in range(a + 1, len(idx)):
                    i, j = idx[a], idx[b]
                    if hamming(hs[i], hs[j]) <= NEAR_DUP_BITS:
                        pairs.add((keys[i], keys[j]))
    return sorted(pairs)


def ddr_table(raw: pd.DataFrame, excl: list) -> pd.DataFrame:
    d = raw[raw.dataset == "DDR"].copy()
    d["role"] = d.source_split.map(ROLE)
    d["grade"] = d.grade_raw
    d["included"] = 1
    for k in d.image_key[d.decode_ok != 1]:
        excl.append((k, "unreadable"))
    for k in d.image_key[(d.decode_ok == 1) & (d.grade == 5)]:
        excl.append((k, "ungradable (grade 5)"))
    ok = d[(d.decode_ok == 1) & (d.grade != 5)]
    for _, g in ok.groupby("pixel_md5"):
        if len(g) < 2:
            continue
        if g.grade.nunique() > 1:
            excl += [(k, "duplicate pixels with conflicting grades") for k in g.image_key]
            continue
        g = g.assign(o=g.role.map(ORDER)).sort_values(["o", "image_key"])
        excl += [(k, "duplicate pixels (copy kept in earliest partition)") for k in g.image_key.iloc[1:]]
    bad = {k for k, _ in excl}
    ok = ok[~ok.image_key.isin(bad)]
    lut = ok.set_index("image_key")
    for a, b in near_dup_pairs(ok):
        ra, rb = lut.loc[a], lut.loc[b]
        if ra.role == rb.role:
            continue
        later = a if ORDER[ra.role] > ORDER[rb.role] else b
        excl.append((later, "near-duplicate of an image in an earlier partition"))
    bad = {k for k, _ in excl}
    d.loc[d.image_key.isin(bad), "included"] = 0
    return d


def attach_masks(d: pd.DataFrame, raw: pd.DataFrame, excl_masks: list) -> pd.DataFrame:
    seg = raw[raw.dataset == "DDRseg"].copy()
    seg["stem"] = seg.image_key.str.replace("ddrseg_", "", regex=False)
    d = d.copy()
    d["stem"] = d.image_key.str.replace("ddr_", "", regex=False)
    d["seg_key"], d["mask_role"], d["seg_source_split"] = "", "", ""
    lut = d.set_index("stem")
    for r in seg.itertuples():
        if r.stem not in lut.index:
            excl_masks.append((r.image_key, "no grading counterpart")); continue
        g = lut.loc[r.stem]
        if g.included != 1:
            excl_masks.append((r.image_key, "grading counterpart excluded")); continue
        if r.pixel_md5 != g.pixel_md5:
            excl_masks.append((r.image_key, "segmentation image differs from grading image")); continue
        if not all(getattr(r, f"mask_shape_ok_{c}", 1) == 1 for c in ["MA", "HE", "EX", "SE"]):
            excl_masks.append((r.image_key, "mask size mismatch")); continue
        i = d.index[d.stem == r.stem][0]
        d.loc[i, ["seg_key", "seg_source_split"]] = [r.image_key, r.source_split]
        for c in ["MA", "HE", "EX", "SE"]:
            d.loc[i, f"mask_{c}"] = getattr(r, f"mask_{c}")
        d.loc[i, "mask_role"] = {"ddr_train": "seg_train", "ddr_dev": "seg_dev", "ddr_test": "seg_test"}[g.role]
    return d.drop(columns="stem")


def messidor2_table(raw: pd.DataFrame, cfg: dict, excl: list) -> pd.DataFrame:
    """Messidor-2：第三外部队列，Krause 等人的五级裁定分级（每张照片）；检查编号取自配对表，重抽样按检查。"""
    lab = pd.read_csv(ROOT / cfg["paths"]["raw_messidor2"] / "labels.csv")
    lab = dict(zip("mes2_" + lab.image.map(lambda v: Path(v).stem), lab.grade))
    g = raw[raw.dataset == "Messidor2"].copy()
    g["grade"] = g.image_key.map(lab)
    g["role"], g["included"] = "messidor2", 1
    pairs = ROOT / cfg["paths"]["raw_messidor2"] / "messidor2_pairs.csv"     # 检查配对：patient_id = 检查编号
    if pairs.exists():
        pr = pd.read_csv(pairs)
        ex = dict(zip(pr.image.str.lower().map(lambda v: "mes2_" + Path(v).stem.lower()), zip(pr.exam, pr.eye)))
        got = g.image_key.str.lower().map(ex)
        assert got.notna().all(), "Messidor-2 photograph without examination"
        g["patient_id"], g["eye"] = got.map(lambda v: v[0]), got.map(lambda v: v[1])
    for k in g.image_key[g.decode_ok != 1]:
        excl.append((k, "unreadable"))
    g.loc[g.decode_ok != 1, "included"] = 0
    return g


def aptos_table(raw: pd.DataFrame, cfg: dict, excl: list) -> pd.DataFrame:
    """APTOS 2019：多相机训练补充集，只进入训练（模型选择仍只用 DDR 开发区）；与任一评价照片近重复者剔除。"""
    lab = pd.read_csv(ROOT / cfg["paths"]["raw_aptos"] / "labels.csv")
    lab = dict(zip("aptos_" + lab.image.map(lambda v: Path(v).stem), lab.grade))
    g = raw[raw.dataset == "APTOS"].copy()
    g["grade"] = g.image_key.map(lab)
    g["role"], g["included"] = "aptos_train", 1
    for k in g.image_key[g.decode_ok != 1]:
        excl.append((k, "unreadable"))
    g.loc[g.decode_ok != 1, "included"] = 0
    return g


def eyepacs_table(raw: pd.DataFrame, cfg: dict, excl: list) -> pd.DataFrame:
    """EyePACS：第四外部队列（Kaggle 2015 训练集，每位患者左右眼各一张照片，五级分级）；按患者重抽样，
    只在全部模型与设置确定后评价一次。"""
    lab = pd.read_csv(ROOT / cfg["paths"]["raw_eyepacs"] / "labels.csv")
    lab["image_key"] = "eyp_" + lab.image.map(lambda v: Path(v).stem)
    g = raw[raw.dataset == "EyePACS"].merge(lab[["image_key", "grade", "patient_id", "eye"]], on="image_key", how="left")
    g["role"], g["included"] = "eyepacs", 1
    for k in g.image_key[g.decode_ok != 1]:
        excl.append((k, "unreadable"))
    g.loc[g.decode_ok != 1, "included"] = 0
    return g


def idrid_table(raw: pd.DataFrame, cfg: dict, excl: list) -> pd.DataFrame:
    """IDRiD：第二外部队列。分级图像（单张照片，无患者编号）按照片评价；分割子集的 81 张照片只用于外部病灶覆盖评价。"""
    base = ROOT / cfg["paths"]["raw_idrid"] / "B. Disease Grading" / "2. Groundtruths"
    lab = {}
    for split, fn in (("train", "a. IDRiD_Disease Grading_Training Labels.csv"), ("test", "b. IDRiD_Disease Grading_Testing Labels.csv")):
        t = pd.read_csv(base / fn)
        for name, g in zip(t["Image name"], t["Retinopathy grade"]):
            lab[f"idrid_{split}_{name}"] = int(g)
    g = raw[raw.dataset == "IDRiD"].copy()
    g["grade"] = g.image_key.map(lab).fillna(-1).astype(int)
    g["role"], g["included"] = "idrid", 1
    for r in g.itertuples():
        if r.decode_ok != 1 or r.grade < 0:
            excl.append((r.image_key, "missing grade or unreadable"))
    g.loc[(g.decode_ok != 1) | (g.grade < 0), "included"] = 0
    sg = raw[raw.dataset == "IDRiDseg"].copy()
    sg["grade"], sg["role"], sg["included"] = -1, "idrid_seg", (sg.decode_ok == 1).astype(int)
    sg["seg_key"], sg["mask_role"], sg["seg_source_split"] = sg.image_key, "ext_seg", sg.source_split
    return pd.concat([g, sg])


def deepdrid_table(raw: pd.DataFrame, cfg: dict, excl: list) -> pd.DataFrame:
    base = ROOT / cfg["paths"]["raw_deepdrid"]
    lab = []
    for f in ["regular-fundus-training/regular-fundus-training.csv", "regular-fundus-validation/regular-fundus-validation.csv"]:
        t = pd.read_csv(base / f)
        lab.append(pd.DataFrame({"image_id": t.image_id, "left": t.left_eye_DR_Level, "right": t.right_eye_DR_Level,
                                 "quality": t["Overall quality"]}))
    lab = pd.concat(lab)
    lab["side"] = lab.image_id.str.split("_").str[1].str[0]
    lab["grade"] = np.where(lab.side == "l", lab.left, lab.right)
    lab["side_conflict"] = ((lab.side == "l") & (lab.left.isna() | lab.right.notna())) | \
                           ((lab.side == "r") & (lab.right.isna() | lab.left.notna()))
    ev = base / "Online-Challenge1&2-Evaluation"
    t1 = list(openpyxl.load_workbook(ev / "Challenge1_labels.xlsx").active.iter_rows(values_only=True))[1:]
    t2 = list(openpyxl.load_workbook(ev / "Challenge2_labels.xlsx").active.iter_rows(values_only=True))[1:]
    q2 = {r[0]: r[1] for r in t2}
    test = pd.DataFrame({"image_id": [r[0] for r in t1], "grade": [r[1] for r in t1]})
    test["quality"] = test.image_id.map(q2); test["side_conflict"] = False
    lab = pd.concat([lab[["image_id", "grade", "quality", "side_conflict"]], test])
    lab["image_key"] = "drid_" + lab.image_id
    d = raw[raw.dataset == "DeepDRiD"].merge(lab, on="image_key", how="left")
    d["patient_id"] = d.image_id.str.split("_").str[0].astype(int)
    d["eye"] = d.image_id.str.split("_").str[1].str[0]
    d["role"], d["included"] = "deepdrid", 1
    badp = set(d.patient_id[d.side_conflict.fillna(False).astype(bool)])
    eye_conf = d.groupby(["patient_id", "eye"]).grade.nunique()
    badp |= {p for (p, _), n in eye_conf.items() if n > 1}
    for r in d.itertuples():
        if r.patient_id in badp:
            excl.append((r.image_key, "eye laterality inconsistent with eye-level labels (whole patient removed)"))
        elif pd.isna(r.grade) or r.decode_ok != 1:
            excl.append((r.image_key, "missing grade or unreadable"))
    d.loc[d.image_key.isin({k for k, _ in excl}), "included"] = 0
    d["grade"] = d.grade.fillna(-1).astype(int)
    return d


def main():
    cfg = load_config()
    raw = pd.read_csv(ROOT / "results/01_preprocess/images_raw.tsv", sep="\t")
    excl, excl_masks = [], []
    ddr = attach_masks(ddr_table(raw, excl), raw, excl_masks)
    drid = deepdrid_table(raw, cfg, excl)
    idr = idrid_table(raw, cfg, excl)
    cross = near_dup_pairs(pd.concat([ddr[ddr.included == 1], drid[drid.included == 1]])[["image_key", "dhash"]])
    cross = [(a, b) for a, b in cross if a[:4] != b[:4]]
    for a, b in cross:
        k = a if a.startswith("drid") else b
        excl.append((k, "near-duplicate of a DDR image"))
        drid.loc[drid.image_key == k, "included"] = 0
    # IDRiD 与开发队列（DDR）之间的近重复：剔除 IDRiD 一方
    cross2 = near_dup_pairs(pd.concat([ddr[ddr.included == 1], idr[idr.included == 1]])[["image_key", "dhash"]])
    cross2 = [(a, b) for a, b in cross2 if a.startswith("ddr_") != b.startswith("ddr_")]
    for a, b in cross2:
        k = b if a.startswith("ddr_") else a
        excl.append((k, "near-duplicate of a DDR image"))
        idr.loc[idr.image_key == k, "included"] = 0
    mes = messidor2_table(raw, cfg, excl)
    cross3 = near_dup_pairs(pd.concat([ddr[ddr.included == 1], mes[mes.included == 1]])[["image_key", "dhash"]])
    cross3 = [(a, b) for a, b in cross3 if a.startswith("ddr_") != b.startswith("ddr_")]
    for a, b in cross3:
        k = b if a.startswith("ddr_") else a
        excl.append((k, "near-duplicate of a DDR image"))
        mes.loc[mes.image_key == k, "included"] = 0
    apt = aptos_table(raw, cfg, excl)
    evals = pd.concat([ddr[(ddr.included == 1) & (ddr.role != "ddr_train")], drid[drid.included == 1],
                       idr[idr.included == 1], mes[mes.included == 1]])
    cross4 = near_dup_pairs(pd.concat([evals, apt[apt.included == 1]])[["image_key", "dhash"]])
    cross4 = [(a, b) for a, b in cross4 if a.startswith("aptos_") != b.startswith("aptos_")]
    for a, b in cross4:
        k = a if a.startswith("aptos_") else b
        excl.append((k, "near-duplicate of an evaluation photograph"))
        apt.loc[apt.image_key == k, "included"] = 0
    # EyePACS：与任一训练或选参照片（DDR 训练/开发区、APTOS）近重复者从 EyePACS 一方剔除；不改动其他队列
    eyp = eyepacs_table(raw, cfg, excl) if (raw.dataset == "EyePACS").any() else raw.iloc[:0].copy()
    train_like = pd.concat([ddr[(ddr.included == 1) & (ddr.role != "ddr_test")], apt[apt.included == 1]])
    cross5 = near_dup_pairs(pd.concat([train_like, eyp[eyp.included == 1]])[["image_key", "dhash"]]) if len(eyp) else []
    cross5 = [(a, b) for a, b in cross5 if a.startswith("eyp_") != b.startswith("eyp_")]
    for a, b in cross5:
        k = a if a.startswith("eyp_") else b
        excl.append((k, "near-duplicate of a training photograph"))
        eyp.loc[eyp.image_key == k, "included"] = 0
    cols = ["image_key", "dataset", "role", "grade", "included", "path", "orig_h", "orig_w", "box_x0", "box_y0",
            "box_side", "fov_frac", "seg_key", "mask_role", "seg_source_split", "patient_id", "eye", "quality",
            "mask_MA", "mask_HE", "mask_EX", "mask_SE"]
    man = pd.concat([ddr, drid, idr, mes, apt, eyp])
    for c in cols:
        if c not in man:
            man[c] = ""
    man = man[cols].copy()
    for c in ["patient_id", "orig_h", "orig_w", "box_x0", "box_y0", "box_side", "quality"]:
        man[c] = man[c].map(lambda v: "" if v == "" or pd.isna(v) else str(int(float(v))))
    out = ROOT / "results/02_manifest"; out.mkdir(parents=True, exist_ok=True)
    man.to_csv(out / "images.tsv", sep="\t", index=False)
    pd.DataFrame(excl, columns=["image_key", "reason"]).drop_duplicates("image_key").to_csv(out / "exclusions.tsv", sep="\t", index=False)
    pd.DataFrame(excl_masks, columns=["seg_key", "reason"]).to_csv(out / "mask_exclusions.tsv", sep="\t", index=False)
    inc = man[man.included == 1]
    summ = {
        "n_by_role": inc.groupby("role").size().to_dict(),
        "grade_by_role": {r: g.grade.value_counts().sort_index().to_dict() for r, g in inc.groupby("role")},
        "masks_by_role": inc[inc.mask_role != ""].groupby("mask_role").size().to_dict(),
        "seg_split_vs_grading_role": {f"{a}->{b}": int(n) for (a, b), n in
                                      inc[inc.mask_role != ""].groupby(["seg_source_split", "role"]).size().items()},
        "deepdrid_patients": int(inc[inc.role == "deepdrid"].patient_id.nunique()),
        "deepdrid_eyes": int(inc[inc.role == "deepdrid"].groupby(["patient_id", "eye"]).ngroups),
        "deepdrid_eye_grades": inc[inc.role == "deepdrid"].groupby(["patient_id", "eye"]).grade.first().value_counts().sort_index().to_dict(),
        "exclusion_reasons": pd.DataFrame(excl, columns=["k", "r"]).drop_duplicates("k").r.value_counts().to_dict(),
        "mask_exclusion_reasons": pd.DataFrame(excl_masks, columns=["k", "r"]).r.value_counts().to_dict(),
        "cross_dataset_near_duplicates": len(cross),
        "idrid_ddr_near_duplicates": len(cross2),
        "idrid_grade_counts": inc[inc.role == "idrid"].grade.astype(int).value_counts().sort_index().to_dict(),
        "messidor2_ddr_near_duplicates": len(cross3),
        "aptos_evaluation_near_duplicates": len(cross4),
        "messidor2_grade_counts": inc[inc.role == "messidor2"].grade.astype(int).value_counts().sort_index().to_dict(),
        "eyepacs_training_near_duplicates": len(cross5),
        "eyepacs_patients": int(inc[inc.role == "eyepacs"].patient_id.nunique()),
        "eyepacs_grade_counts": inc[inc.role == "eyepacs"].grade.astype(int).value_counts().sort_index().to_dict(),
    }
    summ = json.loads(json.dumps(summ, default=lambda o: int(o) if isinstance(o, np.integer) else str(o)))
    (out / "summary.json").write_text(json.dumps(summ, indent=2))
    print(json.dumps(summ, indent=2))
    write_provenance("02_manifest", [ROOT / "results/01_preprocess/images_raw.tsv"], [out / "images.tsv"])


if __name__ == "__main__":
    main()
